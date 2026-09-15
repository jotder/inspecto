#!/usr/bin/env python3
"""
Renders the committed Inspecto Enterprise Whitepaper markdown as a formatted Word document.

RENDERS the .md — it does not re-author it. That distinction is the whole point of this file
(PARAM-free, see BACKLOG `WHITEPAPER-DOCX-DRIFT-1`): until 2026-09-15 the document body was ~12
hand-written blocks of prose that named the markdown only as its OUTPUT path, so the v1.2 rewrite
of the whitepaper never reached the .docx and the generated file still carried v1.1 text. The
.docx is gitignored, so nothing flagged the divergence, and `check-doc-counts` validates only the
.md — meaning every count guard in the repo missed the document a stakeholder actually receives.

Now the .md is the single origin: every guard that covers it covers this output by construction.

The typography, callouts, code blocks, tables, header/footer and page breaks below are the part
that was always worth keeping, and they are unchanged.

Usage:  python scripts/generate_whitepaper_docx.py [output.docx]
"""
import re
import sys
from pathlib import Path
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}>'
                      f'<w:top w:w="{top}" w:type="dxa"/>'
                      f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
                      f'<w:left w:w="{left}" w:type="dxa"/>'
                      f'<w:right w:w="{right}" w:type="dxa"/>'
                      f'</w:tcMar>')
    tcPr.append(tcMar)

def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    tcPr = cell._tc.get_or_add_tcPr()
    borders_elm = parse_xml(f'<w:tcBorders {nsdecls("w")}/>')
    sides = {'top': top, 'bottom': bottom, 'left': left, 'right': right}
    for side, border in sides.items():
        if border:
            val, sz, color = border
            tag = f'<w:{side} {nsdecls("w")} w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            borders_elm.append(parse_xml(tag))
        else:
            tag = f'<w:{side} {nsdecls("w")} w:val="none"/>'
            borders_elm.append(parse_xml(tag))
    tcPr.append(borders_elm)

def format_table(table, col_widths, hdr_bg="1E3A8A", alt_bg="F8FAFC"):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # Format header
    for cell in table.rows[0].cells:
        set_cell_background(cell, hdr_bg)
        set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
        set_cell_borders(cell, bottom=("single", "12", "0F172A"))
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                r.font.name = "Calibri"
                r.font.size = Pt(9.5)

    # Format data rows
    for r_idx, row in enumerate(table.rows[1:]):
        bg = alt_bg if r_idx % 2 == 1 else "FFFFFF"
        for cell in row.cells:
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=90, bottom=90, left=150, right=150)
            set_cell_borders(cell, bottom=("single", "4", "E2E8F0"), top=("none", "0", "auto"))
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.name = "Calibri"
                    r.font.size = Pt(9.0)

    # Set column widths
    for row in table.rows:
        for i, w in enumerate(col_widths):
            row.cells[i].width = Inches(w)

def add_callout(doc, text_paragraphs, border_color="2563EB", bg_color="F1F5F9"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.8)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=180)
    set_cell_borders(cell, left=("single", "32", border_color),
                           top=("none", "0", "auto"),
                           bottom=("none", "0", "auto"),
                           right=("none", "0", "auto"))
    cell.text = ""
    for idx, (p_text, bold_prefix) in enumerate(text_paragraphs):
        p = cell.add_paragraph() if idx > 0 else cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.bold = True
            r_bold.font.name = "Calibri"
            r_bold.font.size = Pt(9.5)
            r_bold.font.color.rgb = RGBColor(15, 23, 42)
        r_body = p.add_run(p_text)
        r_body.font.name = "Calibri"
        r_body.font.size = Pt(9.5)
        r_body.font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph()

def add_code_block(doc, ascii_art):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.8)
    set_cell_background(cell, "0F172A")
    set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
    set_cell_borders(cell, left=("single", "12", "38BDF8"),
                           top=("single", "4", "334155"),
                           bottom=("single", "4", "334155"),
                           right=("single", "4", "334155"))
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(ascii_art.strip("\n"))
    run.font.name = "Consolas"
    run.font.size = Pt(8.0)
    run.font.color.rgb = RGBColor(241, 245, 249)
    doc.add_paragraph()


# ── markdown → the styling helpers above ──────────────────────────────────────
# Deliberately NOT a general markdown implementation. It handles exactly the constructs the
# committed whitepaper uses (measured, not guessed): h1, h3, paragraphs, bullets, ordered lists,
# blockquotes, pipe tables, fenced code, horizontal rules, and inline **bold** / `code`. There are
# no links in the source, so none are rendered. Add a construct when the .md grows one.

INLINE = re.compile(r'(\*\*.+?\*\*|`[^`]+`)')


def _plain(text):
    """Inline markers stripped — for table cells, where format_table styles whole runs."""
    return re.sub(r'\*\*(.+?)\*\*', r'\1', text).replace('`', '')


def _split_row(line):
    return [c.strip() for c in line.strip().strip('|').split('|')]


def _is_divider(line):
    cells = _split_row(line)
    return bool(cells) and all(re.fullmatch(r':?-{2,}:?', c) for c in cells)


def _blocks(md):
    """The .md as a flat list of (kind, payload), so rendering stays a simple dispatch."""
    out, i, lines = [], 0, md.split('\n')
    while i < len(lines):
        line = lines[i]

        if line.startswith('```'):                      # fenced code
            i += 1
            buf = []
            while i < len(lines) and not lines[i].startswith('```'):
                buf.append(lines[i]); i += 1
            i += 1
            out.append(('code', '\n'.join(buf)))
            continue

        if line.startswith('|'):                        # pipe table
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                if not _is_divider(lines[i]):
                    rows.append(_split_row(lines[i]))
                i += 1
            if rows:
                out.append(('table', rows))
            continue

        if line.startswith('> '):                       # blockquote -> callout
            buf = []
            while i < len(lines) and lines[i].startswith('>'):
                buf.append(lines[i].lstrip('>').strip()); i += 1
            out.append(('callout', [p for p in ' '.join(buf).split('\n') if p]))
            continue

        m = re.match(r'^(#{1,6}) +(.*)$', line)
        if m:
            out.append(('h%d' % len(m.group(1)), m.group(2).strip())); i += 1; continue

        # A horizontal rule separates pages in the source; the h1 that follows already starts a
        # new page, so emitting a break here too would leave every page preceded by a blank one.
        if re.fullmatch(r'(---|\*\*\*|___)\s*', line):
            i += 1; continue

        m = re.match(r'^\s*[-*] +(.*)$', line)
        if m:
            out.append(('bullet', m.group(1).strip())); i += 1; continue

        m = re.match(r'^\s*\d+\. +(.*)$', line)
        if m:
            out.append(('ordered', m.group(1).strip())); i += 1; continue

        if line.strip():
            buf = [line.strip()]                        # a paragraph runs to the next blank line
            i += 1
            while i < len(lines) and lines[i].strip() and not re.match(
                    r'^(#{1,6} |\||> |```|\s*[-*] |\s*\d+\. )', lines[i]) and not re.fullmatch(
                    r'(---|\*\*\*|___)\s*', lines[i]):
                buf.append(lines[i].strip()); i += 1
            out.append(('para', ' '.join(buf)))
            continue

        i += 1
    return out


def build_whitepaper_docx(md_path: Path, output_path: Path):
    md = md_path.read_text(encoding='utf-8')
    blocks = _blocks(md)
    doc = Document()

    # Configure Margins
    for s in doc.sections:
        s.top_margin = Inches(0.75)
        s.bottom_margin = Inches(0.75)
        s.left_margin = Inches(0.85)
        s.right_margin = Inches(0.85)

        # Header & Footer
        footer = s.footer
        p_ft = footer.paragraphs[0]
        p_ft.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_ft = p_ft.add_run("Inspecto — Sovereign Data Operations Platform | Confidential")
        r_ft.font.name = "Calibri"
        r_ft.font.size = Pt(8.0)
        r_ft.font.color.rgb = RGBColor(148, 163, 184)

    # Colors Setup
    NAVY = RGBColor(30, 58, 138)       # #1E3A8A
    SLATE = RGBColor(15, 23, 42)       # #0F172A
    MUTED = RGBColor(71, 85, 105)      # #475569
    BLUE = RGBColor(37, 99, 235)       # #2563EB

    # The page count is DERIVED from the source, never hardcoded. The previous generator said
    # "PAGE n OF 10" as a literal while the markdown carried a different number of top-level
    # sections — a second, quieter instance of the same drift this rewrite exists to end.
    total_pages = sum(1 for kind, _ in blocks if kind == 'h1')

    def emit_runs(p, text, size, color=SLATE, bold_all=False):
        """Inline **bold** and `code` as real runs, so emphasis survives into Word."""
        for part in INLINE.split(text):
            if not part:
                continue
            if part.startswith('**') and part.endswith('**') and len(part) > 4:
                r = p.add_run(part[2:-2]); r.bold = True; r.font.name = "Calibri"
            elif part.startswith('`') and part.endswith('`') and len(part) > 2:
                r = p.add_run(part[1:-1]); r.font.name = "Consolas"
                r.font.color.rgb = RGBColor(190, 24, 93)
                r.font.size = Pt(size - 0.5)
                continue
            else:
                r = p.add_run(part); r.font.name = "Calibri"
            r.bold = r.bold or bold_all
            r.font.size = Pt(size)
            r.font.color.rgb = color

    def add_page_header(page_num, title):
        p_pg = doc.add_paragraph()
        p_pg.paragraph_format.space_before = Pt(0)
        p_pg.paragraph_format.space_after = Pt(2)
        r_pg = p_pg.add_run(f"PAGE {page_num} OF {total_pages} • {title.upper()}")
        r_pg.font.name = "Calibri"
        r_pg.font.size = Pt(8.5)
        r_pg.bold = True
        r_pg.font.color.rgb = BLUE

        p_h1 = doc.add_heading(level=1)
        p_h1.paragraph_format.space_before = Pt(0)
        p_h1.paragraph_format.space_after = Pt(4)
        r_h1 = p_h1.add_run(title)
        r_h1.font.name = "Calibri"
        r_h1.font.size = Pt(17)
        r_h1.bold = True
        r_h1.font.color.rgb = NAVY

    def add_h2(title):
        p = doc.add_heading(level=2)
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(title)
        r.font.name = "Calibri"
        r.font.size = Pt(12.5)
        r.bold = True
        r.font.color.rgb = SLATE

    def add_p(text, space_after=6):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        emit_runs(p, text, 10)
        return p

    def add_list_item(text, style):
        p = doc.add_paragraph(style=style)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        emit_runs(p, text, 9.5)

    def add_md_table(rows):
        cols = max(len(r) for r in rows)
        table = doc.add_table(rows=len(rows), cols=cols)
        for r_idx, row in enumerate(rows):
            for c_idx in range(cols):
                table.cell(r_idx, c_idx).text = _plain(row[c_idx]) if c_idx < len(row) else ""
        format_table(table, [round(6.8 / cols, 3)] * cols)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    page = 0
    for kind, payload in blocks:
        if kind == 'h1':
            page += 1
            if page > 1:
                doc.add_page_break()
            add_page_header(page, payload)
        elif kind in ('h2', 'h3', 'h4', 'h5', 'h6'):
            add_h2(payload)
        elif kind == 'para':
            add_p(payload)
        elif kind == 'bullet':
            add_list_item(payload, 'List Bullet')
        elif kind == 'ordered':
            add_list_item(payload, 'List Number')
        elif kind == 'callout':
            # add_callout takes (text, bold_prefix) pairs; the markdown carries no separate prefix.
            add_callout(doc, [(_plain(t), None) for t in payload])
        elif kind == 'code':
            add_code_block(doc, payload)
        elif kind == 'table':
            add_md_table(payload)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return len(blocks), total_pages


# The whitepaper lives beside this script's repo, not at an absolute path on one developer's box —
# the previous generator hardcoded `c:\sandbox\inspecto-clean\...` and could not run on another
# checkout at all.
REPO = Path(__file__).resolve().parent.parent
SOURCE_MD = REPO / "docs" / "stakeholders" / "INSPECTO_ENTERPRISE_WHITEPAPER.md"

if __name__ == "__main__":
    out_file = Path(sys.argv[1]) if len(sys.argv) > 1 else SOURCE_MD.with_suffix(".docx")
    if not SOURCE_MD.exists():
        sys.exit(f"source markdown not found: {SOURCE_MD}")
    blocks, pages = build_whitepaper_docx(SOURCE_MD, out_file)
    print(f"Rendered {SOURCE_MD.name} -> {out_file}")
    print(f"  {blocks} blocks, {pages} pages")
